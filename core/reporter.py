"""
日报生成：docx + html 双格式
"""
from __future__ import annotations

import datetime
from pathlib import Path

from config import platform_report_dir
from core import analyzer
from core.models import ReportData
from core.storage import Storage

PLATFORM_NAMES = {"qq": "QQ", "dingtalk": "钉钉"}


def _today() -> str:
    return datetime.date.today().isoformat()


async def _build_summary(storage: Storage, config: dict, platform: str, date_str: str):
    """收集当日消息 + 重要事项，调用 AI 生成总结"""
    start = int(datetime.datetime.strptime(date_str, "%Y-%m-%d").timestamp() * 1000)
    end = start + 86400_000
    res = storage.query_messages(platform, page=1, page_size=500)
    msgs_raw = [m for m in res["items"] if start <= m["ts"] < end]
    alerts = storage.query_alerts(platform=platform, page=1, page_size=100)["items"]
    today_alerts = [a for a in alerts if start <= a["ts"] < end]

    from core.models import ChatMessage
    msgs = [ChatMessage(msg_id=m["msg_id"], platform=platform, group_id=m["group_id"],
                        group_name=m["group_name"], sender=m["sender"], content=m["content"],
                        msg_type=m["msg_type"], ts=m["ts"]) for m in msgs_raw]
    summary = await analyzer.generate_summary(config, platform, msgs, [])
    return summary, len(msgs), len(today_alerts), msgs_raw, today_alerts


def generate_report(storage: Storage, config: dict, platform: str,
                    date_str: str | None = None, username: str = "") -> ReportData | None:
    """同步入口：生成某日报告（docx + html）；username 指定时输出到用户目录"""
    import asyncio
    date_str = date_str or _today()
    loop = asyncio.new_event_loop()
    try:
        summary, msg_count, imp_count, msgs_raw, alerts = loop.run_until_complete(
            _build_summary(storage, config, platform, date_str))
    finally:
        loop.close()

    if username:
        from config import user_report_dir
        out_dir = user_report_dir(username, platform)
    else:
        out_dir = platform_report_dir(platform)
    out_dir.mkdir(parents=True, exist_ok=True)
    base = out_dir / f"{date_str}-日报"
    docx_path = str(base.with_suffix(".docx"))
    html_path = str(base.with_suffix(".html"))

    _write_docx(docx_path, config, platform, date_str, summary, msgs_raw, alerts)
    _write_html(html_path, config, platform, date_str, summary, msgs_raw, alerts)

    report = ReportData(platform=platform, date=date_str, summary=summary[:2000],
                        docx_path=docx_path, html_path=html_path,
                        msg_count=msg_count, important_count=imp_count)
    storage.save_report(report)
    print(f"[{platform}] 日报已生成: {docx_path} (消息{msg_count} 重要{imp_count})")
    return report


# ── docx ──
def _write_docx(path: str, config: dict, platform: str, date_str: str, summary: str,
                msgs_raw: list[dict], alerts: list[dict]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    name = PLATFORM_NAMES.get(platform, platform.upper())
    doc = Document()
    title = doc.add_heading(f"{name} 聊天分析日报 · {date_str}", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("📊 概要", level=1)
    doc.add_paragraph(f"当日消息数：{len(msgs_raw)} ｜ 重要事项：{len(alerts)}")

    doc.add_heading("📝 AI 总结", level=1)
    for line in summary.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith(("1.", "2.", "3.", "4.", "5.", "- ")):
            doc.add_paragraph(line, style="List Bullet")
        else:
            doc.add_paragraph(line)

    if alerts:
        doc.add_heading("⚠️ 重要事项", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["优先级", "内容", "来源", "执行建议"]):
            hdr[i].text = h
        for a in alerts:
            row = table.add_row().cells
            row[0].text = a["priority"]
            row[1].text = a["content"][:120]
            row[2].text = f"{a['group_name']}/{a['sender']}"
            row[3].text = a["suggestion"][:150]

    doc.add_heading("🗂 消息摘要", level=1)
    for m in msgs_raw[:40]:
        doc.add_paragraph(f"[{m['group_name'] or '群'}][{m['sender']}] {m['content'][:80]}")

    doc.add_paragraph()
    doc.add_paragraph(f"🤖 本报告由智能分析助手自动生成（{config.get('AI_MODEL', 'deepseek-v4-flash-0731')}）")
    doc.save(path)


# ── html ──
def _esc(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _write_html(path: str, config: dict, platform: str, date_str: str, summary: str,
                msgs_raw: list[dict], alerts: list[dict]) -> None:
    model = config.get("AI_MODEL", "deepseek-v4-flash-0731")
    name = PLATFORM_NAMES.get(platform, platform.upper())
    alert_rows = "".join(
        f'<tr><td class="p-{a["priority"]}">{a["priority"]}</td><td>{_esc(a["content"])}</td>'
        f'<td>{_esc(a["group_name"])}/{_esc(a["sender"])}</td><td>{_esc(a["suggestion"])}</td></tr>'
        for a in alerts
    ) or '<tr><td colspan="4">今日无重要事项</td></tr>'
    msg_items = "".join(
        f'<li><b>{_esc(m["sender"])}</b> <span class="grp">{_esc(m["group_name"])}</span>'
        f'<div class="c">{_esc(m["content"])}</div></li>'
        for m in msgs_raw[:60]
    ) or "<li>无消息</li>"
    summary_html = "".join(
        f"<p>{_esc(line)}</p>" for line in summary.split("\n") if line.strip()
    )
    html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{name}日报 {date_str}</title>
<style>
body{{font-family:system-ui,-apple-system,'Microsoft YaHei',sans-serif;background:#0f1115;color:#e6e6e6;margin:0;padding:24px}}
.wrap{{max-width:860px;margin:0 auto}}
h1{{text-align:center;color:#fff}}
h2{{color:#7cc4ff;border-left:4px solid #4a9eff;padding-left:10px;margin-top:32px}}
.card{{background:#181c24;border:1px solid #2a3040;border-radius:10px;padding:16px 20px;margin:12px 0}}
table{{width:100%;border-collapse:collapse}}
th,td{{border:1px solid #2a3040;padding:8px 10px;text-align:left;font-size:14px}}
th{{background:#222834}}
.p-high{{color:#ff6b6b;font-weight:bold}}.p-medium{{color:#ffa94d;font-weight:bold}}.p-low{{color:#69db7c}}
ul{{padding-left:20px}}li{{margin:8px 0}}.grp{{color:#8a94a6;font-size:12px;margin-left:8px}}.c{{color:#c8cedb;margin-top:2px}}
.foot{{text-align:center;color:#6b7280;font-size:12px;margin-top:40px}}
</style></head><body><div class="wrap">
<h1>{name} 聊天分析日报 · {date_str}</h1>
<div class="card"><b>📊 概要</b>：当日消息 {len(msgs_raw)} 条 ｜ 重要事项 {len(alerts)} 条</div>
<h2>📝 AI 总结</h2><div class="card">{summary_html}</div>
<h2>⚠️ 重要事项</h2><div class="card"><table><tr><th>优先级</th><th>内容</th><th>来源</th><th>执行建议</th></tr>{alert_rows}</table></div>
<h2>🗂 消息摘要</h2><div class="card"><ul>{msg_items}</ul></div>
<div class="foot">🤖 本报告由智能分析助手自动生成（{model}）</div>
</div></body></html>"""
    Path(path).write_text(html, encoding="utf-8")
